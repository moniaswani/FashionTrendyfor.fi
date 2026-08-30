"""
Loads runway_metadata_schema.xlsx (Looks + Items sheets) into the
RunwayLooks / RunwayItems DynamoDB tables. See src/dynamodb_setup.py
for table creation and dynamodb_schema notes in that file's docstring.

Excel schema (see the "Schema - Looks" / "Schema - Items" sheets):
    Looks: look_id, designer, designer_lower, season, season_lower,
           collection, event, city, look_number, publication_date,
           source_url, color_palette_hex, mood_keywords, scraped_at,
           silhouette, materials_note, movement, styling
    Items: look_id (FK -> Looks.look_id), item_name, color, material,
           item_label

Items are denormalized with designer/designer_lower/season/season_lower
from their parent look, so RunwayItems can be queried directly via its
DesignerSeasonIndex GSI without a join back to RunwayLooks.

Conversely, each look is written with item_names/colors/materials
aggregated from its own items (see derive_look_facets()), so RunwayLooks
rows are filterable on those facets without joining RunwayItems back.
Existing rows loaded before this existed need src/backfill_runway_looks_facets.py.

Usage:
    python src/excel_to_dynamodb.py runway_metadata_schema.xlsx --dry-run
    python src/excel_to_dynamodb.py runway_metadata_schema.xlsx
"""

import argparse
import os
import sys

import boto3
import openpyxl

REGION = os.environ.get("AWS_REGION", "eu-west-2")
LOOKS_TABLE = "RunwayLooks"
ITEMS_TABLE = "RunwayItems"

LOOKS_LIST_FIELDS = {"color_palette_hex", "mood_keywords"}

# Some shows get published under a shortened/alternate brand name at the
# source (e.g. NOWFASHION's Acne Spring Summer 2024 page is titled "Acne"
# while every other Acne Studios season is titled "Acne Studios"). Since
# designer_lower is the DesignerSeasonIndex partition key, an unnormalized
# alias silently splits one designer's looks across two GSI partitions.
# Keyed by the raw designer name lowercased; value is the canonical
# as-published display name to use instead.
DESIGNER_ALIASES = {
    "acne": "Acne Studios",
    "chloé": "Chloe",
    "catwalkpictures": "Ganni",  # ganni-rtw-ss26-paris scraped a photo-credit line as designer
}


def _clean(v):
    """Blank string -> None, otherwise strip strings."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def normalize_designer(name: str) -> str:
    """Maps known aliases to their canonical display name (see DESIGNER_ALIASES).
    Unrecognized names pass through unchanged, preserving as-published casing."""
    return DESIGNER_ALIASES.get(name.lower(), name)


def read_sheet(ws) -> list[dict]:
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    return [dict(zip(header, row)) for row in rows]


def validate_look(row: dict) -> str | None:
    """Returns an error string, or None if valid."""
    if not _clean(row.get("look_id")):
        return "missing look_id"
    if not _clean(row.get("designer_lower")):
        return "missing designer_lower"
    if not _clean(row.get("season_lower")):
        return "missing season_lower"
    return None


def transform_look(row: dict) -> dict:
    item = {"look_id": _clean(row["look_id"])}
    for field in (
        "designer", "designer_lower", "season", "season_lower",
        "collection", "event", "city", "publication_date",
        "source_url", "scraped_at", "silhouette", "materials_note",
        "movement", "styling",
    ):
        v = _clean(row.get(field))
        if v is not None:
            item[field] = v

    if "designer" in item:
        # designer_lower is derived fresh from the (alias-corrected) designer
        # name rather than trusted from the sheet, since the sheet's own
        # designer_lower is just a lowercased mirror of its designer column
        # and would carry the same alias inconsistency forward otherwise.
        item["designer"] = normalize_designer(item["designer"])
        item["designer_lower"] = item["designer"].lower()

    look_number = _clean(row.get("look_number"))
    if look_number is not None:
        try:
            item["look_number"] = int(look_number)
        except (TypeError, ValueError):
            pass

    for field in LOOKS_LIST_FIELDS:
        v = _clean(row.get(field))
        if v is not None:
            item[field] = [part.strip() for part in v.split(";") if part.strip()]

    return item


def derive_look_facets(item_rows: list[dict]) -> dict:
    """Aggregates item_name/color/material across a look's items onto the
    look itself, mirroring the join ui_prototypes/runway-gallery/scripts/
    export_data.py does when building index.json — denormalized here so
    the regenerateRunwayCache Lambda can return facet-filterable looks
    without joining RunwayItems per look."""
    item_names = sorted({v for r in item_rows if (v := _clean(r.get("item_name")))})
    colors = sorted({v for r in item_rows if (v := _clean(r.get("color")))})
    materials = sorted({v for r in item_rows if (v := _clean(r.get("material")))})
    return {"item_names": item_names, "colors": colors, "materials": materials}


def transform_items(look_id: str, look: dict, item_rows: list[dict]) -> list[dict]:
    out = []
    for idx, row in enumerate(item_rows):
        item = {
            "look_id": look_id,
            "item_index": idx,
        }
        for field in ("item_name", "color", "material", "item_label"):
            v = _clean(row.get(field))
            if v is not None:
                item[field] = v
        # Denormalize from parent look so RunwayItems can be queried
        # standalone via DesignerSeasonIndex.
        for field in ("designer", "designer_lower", "season", "season_lower"):
            v = look.get(field)
            if v is not None:
                item[field] = v
        out.append(item)
    return out


def load_workbook_data(xlsx_path: str):
    """Returns (looks_rows, items_by_look_id)."""
    wb_looks = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    looks_rows = read_sheet(wb_looks["Looks"])
    wb_looks.close()

    wb_items = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    items_rows = read_sheet(wb_items["Items"])
    wb_items.close()

    items_by_look_id: dict[str, list[dict]] = {}
    for row in items_rows:
        lid = _clean(row.get("look_id"))
        if lid is None:
            continue
        items_by_look_id.setdefault(lid, []).append(row)

    return looks_rows, items_by_look_id


def batch_ingest(xlsx_path: str, dry_run: bool = False) -> dict:
    looks_rows, items_by_look_id = load_workbook_data(xlsx_path)

    report = {
        "looks_total": len(looks_rows),
        "looks_ingested": 0,
        "looks_invalid": 0,
        "items_total": 0,
        "items_ingested": 0,
        "errors": [],
    }

    seen_look_ids = set()

    def _run(looks_writer=None, items_writer=None):
        for row in looks_rows:
            err = validate_look(row)
            if err:
                report["looks_invalid"] += 1
                report["errors"].append(f"{row.get('look_id')!r}: {err}")
                continue

            look_id = _clean(row["look_id"])
            if look_id in seen_look_ids:
                report["looks_invalid"] += 1
                report["errors"].append(f"{look_id!r}: duplicate look_id")
                continue
            seen_look_ids.add(look_id)

            look_item = transform_look(row)
            item_rows = items_by_look_id.get(look_id, [])
            look_item.update(derive_look_facets(item_rows))
            garment_items = transform_items(look_id, look_item, item_rows)

            report["looks_ingested"] += 1
            report["items_total"] += len(garment_items)
            report["items_ingested"] += len(garment_items)

            if looks_writer is not None:
                looks_writer.put_item(Item=look_item)
                for gi in garment_items:
                    items_writer.put_item(Item=gi)

    if dry_run:
        _run()
    else:
        dynamodb = boto3.resource("dynamodb", region_name=REGION)
        looks_table = dynamodb.Table(LOOKS_TABLE)
        items_table = dynamodb.Table(ITEMS_TABLE)
        # Nested context managers so both writers flush their final
        # (sub-25-item) batch on exit — BatchWriter has no public close().
        with looks_table.batch_writer() as looks_writer:
            with items_table.batch_writer() as items_writer:
                _run(looks_writer, items_writer)

    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", help="Path to runway_metadata_schema.xlsx")
    parser.add_argument("--dry-run", action="store_true", help="Parse & validate only, no writes")
    args = parser.parse_args()

    if not os.path.exists(args.xlsx_path):
        print(f"❌ File not found: {args.xlsx_path}")
        sys.exit(1)

    mode = "DRY RUN" if args.dry_run else f"LIVE (writing to {LOOKS_TABLE} / {ITEMS_TABLE} in {REGION})"
    print(f"Loading {args.xlsx_path} [{mode}]...")

    report = batch_ingest(args.xlsx_path, dry_run=args.dry_run)

    print()
    print(f"Looks:  {report['looks_ingested']}/{report['looks_total']} ingested"
          f" ({report['looks_invalid']} invalid)")
    print(f"Items:  {report['items_ingested']}/{report['items_total']} ingested")
    if report["errors"]:
        print(f"\nFirst {min(10, len(report['errors']))} errors:")
        for e in report["errors"][:10]:
            print(f"  - {e}")


if __name__ == "__main__":
    main()
